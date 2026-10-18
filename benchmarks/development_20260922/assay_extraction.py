import sys,shutil,subprocess,csv,collections,json
from pathlib import Path
from ppiblocker.common import Run,record_file,atomic_json,checked_command,read_json
root=Path("/lustre/users/u21111210118/projects/test9");src=root/"steps/47_publisher_supplements/runs/20260922T103153Z_d35874272d99/outputs"
run=Run(root,"49_assay_extraction",[record_file(src/"receipts.json")])
shutil.copy2(__file__,run.path/"qc/driver.py")
env=dict(__import__("os").environ);env["PYTHONNOUSERSITE"]="1";env.pop("PYTHONPATH",None)
checked_command([root/"env/metrics/bin/python","-m","pip","install","pypdf","openpyxl"],run.path/"work",run.path/"logs/install.log",env,timeout=600)
worker=run.path/"work/extract.py"
worker.write_text("""
import sys,csv,json,collections
from pathlib import Path
from pypdf import PdfReader
from openpyxl import load_workbook
src=Path(sys.argv[1]);out=Path(sys.argv[2])
for p in src.glob("*.pdf"):
 text="\\n".join(f"PAGE {i+1}\\n"+(x.extract_text() or "") for i,x in enumerate(PdfReader(p).pages))
 (out/(p.stem+".txt")).write_text(text)
 print("PDF",p.name,"pages",len(PdfReader(p).pages),"chars",len(text),flush=True)
for p in src.glob("*.csv"):
 with p.open() as f:r=csv.DictReader(f);rows=list(r);cols=r.fieldnames
 print("CSV",p.name,len(rows),cols,flush=True)
 (out/(p.stem+"_summary.json")).write_text(json.dumps(dict(rows=len(rows),columns=cols,example=rows[:2]),indent=2))
for p in src.glob("*.xlsx"):
 w=load_workbook(p,read_only=True,data_only=True);summary={}
 for s in w:
  rows=list(s.values);summary[s.title]=dict(rows=s.max_row,cols=s.max_column,preview=rows[:4])
 print("XLSX",p.name,json.dumps(summary,default=str),flush=True)
 (out/(p.stem+"_summary.json")).write_text(json.dumps(summary,default=str,indent=2))
""")
checked_command([root/"env/metrics/bin/python",worker,src,run.path/"outputs"],run.path/"work",run.path/"logs/extract.log",env,timeout=600)
for p in run.path.joinpath("outputs").iterdir():run.output(p.name,p)
print(run.finish("PRIMARY_ASSAY_TEXT_EXTRACTED","primary_source_not_label_adjudication"),flush=True)
